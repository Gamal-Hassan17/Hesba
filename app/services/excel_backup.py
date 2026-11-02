
from io import BytesIO
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models import (
    FinancialAccount,
    Expense,
    FundingTransaction,
    RecurringExpense,
    User,
)

FILENAME = "Hesba_Backup.xlsx"
VERSION = 4

# العمود الأول يحتفظ بالمفتاح الداخلي للمزامنة، وسيتم إخفاؤه.
TABLES = {
    "البنود_المالية": (
        [
            "مفتاح_مزامنة",
            "اسم_البند",
            "الوصف",
            "الميزانية_الأولية",
            "إجمالي_الأموال_المضافة",
            "إجمالي_المصروفات",
            "الرصيد_المتبقي",
            "له_ميزانية",
            "مؤرشف",
            "تاريخ_الإنشاء",
            "تاريخ_التحديث",
        ],
        FinancialAccount,
    ),
    "تفاصيل_المصروفات": (
        [
            "مفتاح_مزامنة",
            "اسم_البند",
            "المبلغ",
            "وصف_المصروف",
            "تاريخ_المصروف",
            "تاريخ_الإضافة",
            "تاريخ_التحديث",
            "ملف_الإيصال",
        ],
        Expense,
    ),
    "تفاصيل_الأموال_المضافة": (
        [
            "مفتاح_مزامنة",
            "اسم_البند",
            "المبلغ",
            "الوصف",
            "تاريخ_العملية",
            "تاريخ_الإضافة",
        ],
        FundingTransaction,
    ),
    "المصروفات_المتكررة": (
        [
            "مفتاح_مزامنة",
            "اسم_البند",
            "المبلغ",
            "الوصف",
            "معدل_التكرار",
            "موعد_الاستحقاق_القادم",
            "تاريخ_النهاية",
            "نشط",
            "تاريخ_الإنشاء",
        ],
        RecurringExpense,
    ),
}

META_SHEET = "System_Metadata"

META_HEADERS = [
    "User_ID",
    "Backup_ID",
    "Backup_Version",
    "Last_Sync_ID",
    "Created_At",
]

LEGACY_SHEETS = {
    "Accounts": "البنود_المالية",
    "Expenses": "تفاصيل_المصروفات",
    "Funding": "تفاصيل_الأموال_المضافة",
}

LEGACY_ACCOUNTS_V1 = [
    "ID",
    "Name",
    "Description",
    "Initial_Amount",
    "Is_Budgeted",
    "Is_Archived",
    "Created_At",
    "Updated_At",
]

LEGACY_HEADERS = {
    "البنود_المالية": [
        "ID", "Name", "Description", "Initial_Amount",
        "Total_Funding", "Total_Expenses", "Remaining_Balance",
        "Is_Budgeted", "Is_Archived", "Created_At", "Updated_At",
    ],
    "تفاصيل_المصروفات": [
        "ID", "Account_ID", "Amount", "Description",
        "Expense_Date", "Created_At", "Updated_At",
    ],
    "تفاصيل_الأموال_المضافة": [
        "ID", "Account_ID", "Amount", "Description",
        "Transaction_Date", "Created_At",
    ],
}

# رؤوس الإصدار السابق قبل استبدال رقم البند باسمه.
OLD_HEADERS = {
    "البنود_المالية": [
        "رقم_البند", "اسم_البند", "الوصف",
        "الميزانية_الأولية", "إجمالي_الأموال_المضافة",
        "إجمالي_المصروفات", "الرصيد_المتبقي",
        "له_ميزانية", "مؤرشف", "تاريخ_الإنشاء",
        "تاريخ_التحديث",
    ],
    "تفاصيل_المصروفات": [
        "رقم_المصروف", "رقم_البند", "المبلغ",
        "وصف_المصروف", "تاريخ_المصروف",
        "تاريخ_الإضافة", "تاريخ_التحديث", "ملف_الإيصال",
    ],
    "تفاصيل_الأموال_المضافة": [
        "رقم_العملية", "رقم_البند", "المبلغ",
        "الوصف", "تاريخ_العملية", "تاريخ_الإضافة",
    ],
    "المصروفات_المتكررة": [
        "رقم_القاعدة", "رقم_البند", "المبلغ",
        "الوصف", "معدل_التكرار",
        "موعد_الاستحقاق_القادم", "تاريخ_النهاية",
        "نشط", "تاريخ_الإنشاء",
    ],
}


def _value(value):
    if value is None:
        return None

    if hasattr(value, "isoformat"):
        return (
            value.isoformat(sep=" ")
            if isinstance(value, datetime)
            else value.isoformat()
        )

    if hasattr(value, "as_tuple"):
        return float(value)

    return value


def _account_names(user_id):
    accounts = FinancialAccount.query.filter_by(
        user_id=user_id
    ).all()

    return {
        str(account.id): account.name
        for account in accounts
    }


def _current_rows(model, columns, user_id):
    query = model.query

    if model is FinancialAccount:
        query = query.filter(
            FinancialAccount.user_id == user_id
        )
    else:
        query = query.join(
            FinancialAccount,
            model.account_id == FinancialAccount.id
        ).filter(
            FinancialAccount.user_id == user_id
        )

    return [
        [
            _value(getattr(obj, attr))
            for attr in columns
        ]
        for obj in query.order_by(model.id.asc()).all()
    ]


def _style(ws):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="166534"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    for col in ws.columns:
        letter = get_column_letter(col[0].column)

        width = min(
            max(
                max(
                    len(str(c.value or ""))
                    for c in col
                ) + 2,
                12,
            ),
            40,
        )

        ws.column_dimensions[letter].width = width

    # إخفاء مفتاح المزامنة الداخلي.
    ws.column_dimensions["A"].hidden = True


def _read_metadata(meta_ws):
    first_row = [
        cell.value
        for cell in meta_ws[1]
    ]

    if "User_ID" in first_row:
        values = (
            [cell.value for cell in meta_ws[2]]
            if meta_ws.max_row >= 2
            else []
        )

        return {
            str(header): (
                values[index]
                if index < len(values)
                else None
            )
            for index, header in enumerate(first_row)
            if header is not None
        }

    return {
        str(row[0].value): row[1].value
        for row in meta_ws.iter_rows(
            min_row=1,
            max_col=2
        )
        if row[0].value is not None
    }


def _prepare_sheet(ws, name, headers, old_version=None):
    actual = [
        cell.value
        for cell in ws[1]
    ]

    if (
        name == "البنود_المالية"
        and old_version == 1
        and actual == LEGACY_ACCOUNTS_V1
    ):
        legacy_rows = list(
            ws.iter_rows(
                min_row=2,
                values_only=True
            )
        )

        ws.delete_rows(1, ws.max_row)
        ws.append(headers)

        for row in legacy_rows:
            row = list(row)

            ws.append(
                row[:4]
                + [None, None, None]
                + row[4:8]
            )

        return

    if actual == headers:
        return

    if actual == OLD_HEADERS.get(name):
        for col, header in enumerate(headers, start=1):
            ws.cell(
                row=1,
                column=col,
                value=header
            )

        return

    if actual == LEGACY_HEADERS.get(name):
        for col, header in enumerate(headers, start=1):
            ws.cell(
                row=1,
                column=col,
                value=header
            )

        if name == "تفاصيل_المصروفات":
            if ws.max_column < len(headers):
                ws.cell(
                    row=1,
                    column=len(headers),
                    value=headers[-1]
                )

        return

    raise ValueError(
        f"تنسيق ورقة {name} غير متوافق؛ لم يتم تعديل الملف."
    )


def _replace_account_ids(ws, account_names):
    """
    تحويل أرقام البنود القديمة إلى أسماء البنود.
    يتم ذلك في العمود B فقط.
    """

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(
            row=row,
            column=2
        )

        value = cell.value

        if value is None:
            continue

        key = str(value)

        if key in account_names:
            cell.value = account_names[key]


def _upsert_rows(ws, current_data):
    index = {
        str(ws.cell(row=row, column=1).value): row
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=1).value is not None
    }

    for values in current_data:
        key = str(values[0])

        if key in index:
            row_num = index[key]

            for col_num, value in enumerate(
                values,
                start=1
            ):
                ws.cell(
                    row=row_num,
                    column=col_num,
                    value=value
                )
        else:
            ws.append(values)


def merge_workbook(uploaded_bytes, user_id):
    """Create or update a user's complete Excel backup."""

    now = datetime.utcnow().isoformat(sep=" ")

    sync_id = (
        f"sync-{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}"
    )

    current_user = User.query.get(user_id)

    if not current_user:
        raise ValueError(
            "تعذر التحقق من صاحب النسخة الاحتياطية."
        )

    account_names = _account_names(user_id)

    if uploaded_bytes:
        try:
            wb = load_workbook(
                BytesIO(uploaded_bytes)
            )
        except Exception as exc:
            raise ValueError(
                "ملف Excel غير صالح أو تالف؛ لم يتم تغييره."
            ) from exc

        if META_SHEET not in wb.sheetnames:
            raise ValueError(
                "الملف لا يحتوي على بيانات تعريف النسخة؛ "
                "أوقفنا التحديث لحماية بياناتك."
            )

        meta = _read_metadata(wb[META_SHEET])

        try:
            same_owner = (
                int(float(meta.get("User_ID")))
                == int(user_id)
            )
        except (TypeError, ValueError, OverflowError):
            same_owner = (
                str(meta.get("User_ID", "")).strip()
                == str(user_id)
            )

        if not same_owner:
            raise ValueError(
                "هذه النسخة مرتبطة بحساب مختلف. لم يتم تعديلها."
            )

        try:
            old_version = int(
                meta.get("Backup_Version") or 0
            )
        except (TypeError, ValueError):
            old_version = 0

        if old_version not in (1, 2, 3, VERSION):
            raise ValueError(
                "إصدار ملف Excel غير متوافق. لم يتم تعديل الملف."
            )

        # تحويل أسماء الأوراق القديمة.
        for old_name, new_name in LEGACY_SHEETS.items():
            if (
                old_name in wb.sheetnames
                and new_name not in wb.sheetnames
            ):
                wb[old_name].title = new_name

        for name, (headers, _) in TABLES.items():
            if name not in wb.sheetnames:
                wb.create_sheet(name).append(headers)
                continue

            _prepare_sheet(
                wb[name],
                name,
                headers,
                old_version
            )

        backup_id = meta.get("Backup_ID")
        created_at = meta.get("Created_At") or now

    else:
        wb = Workbook()
        wb.remove(wb.active)

        backup_id = (
            f"hesba-{user_id}-"
            f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"
        )

        created_at = now

        for name, (headers, _) in TABLES.items():
            wb.create_sheet(name).append(headers)

        wb.create_sheet(
            META_SHEET
        ).append(META_HEADERS)

    # تحديث بيانات كل ورقة.
    for name, (headers, model) in TABLES.items():
        ws = wb[name]

        attrs = {
            "تفاصيل_المصروفات": [
                "id",
                "account_id",
                "amount",
                "description",
                "expense_date",
                "created_at",
                "updated_at",
                "receipt_path",
            ],
            "تفاصيل_الأموال_المضافة": [
                "id",
                "account_id",
                "amount",
                "description",
                "transaction_date",
                "created_at",
            ],
            "المصروفات_المتكررة": [
                "id",
                "account_id",
                "amount",
                "description",
                "frequency",
                "next_due",
                "end_date",
                "is_active",
                "created_at",
            ],
        }.get(name)

        if name == "البنود_المالية":
            current_data = []

            accounts = (
                model.query
                .filter_by(user_id=user_id)
                .order_by(model.id.asc())
                .all()
            )

            for account in accounts:
                current_data.append([
                    _value(account.id),
                    _value(account.name),
                    _value(account.description),
                    _value(account.initial_amount),
                    _value(account.total_funding),
                    _value(account.total_expenses),
                    (
                        _value(account.remaining_balance)
                        if account.is_budgeted
                        else None
                    ),
                    _value(account.is_budgeted),
                    _value(account.is_archived),
                    _value(account.created_at),
                    _value(account.updated_at),
                ])

        else:
            current_data = _current_rows(
                model,
                attrs,
                user_id
            )

            # استبدال account_id باسم البند.
            for values in current_data:
                values[1] = account_names.get(
                    str(values[1]),
                    "بند محذوف"
                )

            # تحديث الأسماء الموجودة في النسخ القديمة.
            _replace_account_ids(
                ws,
                account_names
            )

        _upsert_rows(
            ws,
            current_data
        )

        _style(ws)

    # تحديث بيانات النسخة الاحتياطية.
    meta_ws = wb[META_SHEET]

    metadata = {
        "User_ID": user_id,
        "Backup_ID": backup_id,
        "Backup_Version": VERSION,
        "Last_Sync_ID": sync_id,
        "Created_At": created_at,
    }

    meta_ws.delete_rows(
        1,
        meta_ws.max_row
    )

    meta_ws.append(META_HEADERS)

    meta_ws.append([
        metadata[key]
        for key in META_HEADERS
    ])

    meta_ws.sheet_state = "hidden"

    output = BytesIO()

    wb.save(output)
    output.seek(0)

    return output.getvalue()


def merge_account_workbook(
    uploaded_bytes,
    user_id,
    account_id
):
    """Create/update a backup for one financial account."""

    account = FinancialAccount.query.filter_by(
        id=account_id,
        user_id=user_id
    ).first()

    if not account:
        raise ValueError(
            "البند المالي غير موجود أو لا يتبع حسابك."
        )

    user = User.query.get(user_id)

    if not user:
        raise ValueError(
            "تعذر التحقق من صاحب النسخة الاحتياطية."
        )

    now = datetime.utcnow().isoformat(sep=" ")

    sheets = {
        "ملخص_البند": [
            "مفتاح_مزامنة",
            "اسم_البند",
            "الوصف",
            "الميزانية_الأولية",
            "إجمالي_الأموال_المضافة",
            "إجمالي_المصروفات",
            "الرصيد_المتبقي",
            "له_ميزانية",
        ],

        "تفاصيل_المصروفات": TABLES[
            "تفاصيل_المصروفات"
        ][0],

        "تفاصيل_الأموال_المضافة": TABLES[
            "تفاصيل_الأموال_المضافة"
        ][0],

        "المصروفات_المتكررة": TABLES[
            "المصروفات_المتكررة"
        ][0],
    }

    if uploaded_bytes:
        try:
            wb = load_workbook(
                BytesIO(uploaded_bytes)
            )
        except Exception as exc:
            raise ValueError(
                "ملف Excel غير صالح أو تالف."
            ) from exc

        if META_SHEET not in wb.sheetnames:
            raise ValueError(
                "الملف لا يحتوي على بيانات تعريف؛ "
                "اختر ملف البند الذي أنشأه حِسبة."
            )

        meta = _read_metadata(
            wb[META_SHEET]
        )

        try:
            if (
                int(float(meta.get("User_ID"))) != int(user_id)
                or int(float(meta.get("Account_ID"))) != int(account_id)
            ):
                raise ValueError(
                    "الملف مرتبط بحساب مستخدم أو بند مختلف."
                )

        except (TypeError, ValueError, OverflowError) as exc:
            if (
                isinstance(exc, ValueError)
                and str(exc) ==
                "الملف مرتبط بحساب مستخدم أو بند مختلف."
            ):
                raise

            raise ValueError(
                "تعذر التحقق من ملكية ملف Excel."
            ) from exc

        for name, headers in sheets.items():
            if name not in wb.sheetnames:
                raise ValueError(
                    f"ورقة {name} غير موجودة في الملف."
                )

            ws = wb[name]

            actual = [
                cell.value
                for cell in ws[1]
            ]

            old_headers = (
                OLD_HEADERS.get(name)
                if name != "ملخص_البند"
                else [
                    "رقم_البند",
                    "اسم_البند",
                    "الوصف",
                    "الميزانية_الأولية",
                    "إجمالي_الأموال_المضافة",
                    "إجمالي_المصروفات",
                    "الرصيد_المتبقي",
                    "له_ميزانية",
                ]
            )

            if actual == headers:
                continue

            if actual == old_headers:
                for col, header in enumerate(
                    headers,
                    start=1
                ):
                    ws.cell(
                        row=1,
                        column=col,
                        value=header
                    )
            else:
                raise ValueError(
                    f"تنسيق ورقة {name} غير متوافق."
                )

    else:
        wb = Workbook()
        wb.remove(wb.active)

        for name, headers in sheets.items():
            wb.create_sheet(name).append(headers)

        wb.create_sheet(
            META_SHEET
        ).append([
            "User_ID",
            "Account_ID",
            "Backup_Version",
            "Created_At",
        ])

    # تحديث ملخص البند.
    summary = wb["ملخص_البند"]

    summary_values = [
        account.id,
        account.name,
        account.description,
        _value(account.initial_amount),
        _value(account.total_funding),
        _value(account.total_expenses),
        (
            _value(account.remaining_balance)
            if account.is_budgeted
            else None
        ),
        _value(account.is_budgeted),
    ]

    if summary.max_row >= 2:
        for col, value in enumerate(
            summary_values,
            start=1
        ):
            summary.cell(
                row=2,
                column=col,
                value=value
            )
    else:
        summary.append(summary_values)

    datasets = {
        "تفاصيل_المصروفات": (
            Expense,
            [
                "id",
                "account_id",
                "amount",
                "description",
                "expense_date",
                "created_at",
                "updated_at",
                "receipt_path",
            ],
        ),

        "تفاصيل_الأموال_المضافة": (
            FundingTransaction,
            [
                "id",
                "account_id",
                "amount",
                "description",
                "transaction_date",
                "created_at",
            ],
        ),

        "المصروفات_المتكررة": (
            RecurringExpense,
            [
                "id",
                "account_id",
                "amount",
                "description",
                "frequency",
                "next_due",
                "end_date",
                "is_active",
                "created_at",
            ],
        ),
    }

    for sheet_name, (model, attrs) in datasets.items():
        ws = wb[sheet_name]

        # في نسخة البند الواحد، يظهر اسم البند
        # بدل رقم البند في جميع السجلات.
        for row in range(2, ws.max_row + 1):
            ws.cell(
                row=row,
                column=2,
                value=account.name
            )

        query = (
            model.query
            .filter_by(account_id=account.id)
            .order_by(model.id.asc())
        )

        current_data = []

        for obj in query.all():
            values = [
                _value(getattr(obj, attr))
                for attr in attrs
            ]

            values[1] = account.name

            current_data.append(values)

        _upsert_rows(
            ws,
            current_data
        )

        _style(ws)

    _style(summary)

    # بيانات تعريف النسخة.
    meta_ws = wb[META_SHEET]

    created_at = now

    if uploaded_bytes:
        previous = _read_metadata(meta_ws)
        created_at = previous.get("Created_At") or now

    meta_ws.delete_rows(
        1,
        meta_ws.max_row
    )

    meta_ws.append([
        "User_ID",
        "Account_ID",
        "Backup_Version",
        "Created_At",
    ])

    meta_ws.append([
        user_id,
        account_id,
        VERSION,
        created_at,
    ])

    meta_ws.sheet_state = "hidden"

    output = BytesIO()

    wb.save(output)
    output.seek(0)

    return output.getvalue()
