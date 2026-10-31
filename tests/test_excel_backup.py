from io import BytesIO
from openpyxl import load_workbook
from app import db
from app.models import User, FinancialAccount, Expense
from app.services.excel_backup import merge_workbook

def make_user():
    user = User(username="owner")
    user.set_password("strong-password")
    db.session.add(user)
    db.session.commit()
    return user

def test_excel_backup_creates_workbook_and_metadata(app):
    with app.app_context():
        user = make_user()
        db.session.add(FinancialAccount(name="Home", initial_amount=100))
        db.session.commit()
        content = merge_workbook(None, user.id)
        wb = load_workbook(BytesIO(content))
        assert {"البنود_المالية", "تفاصيل_المصروفات", "تفاصيل_الأموال_المضافة", "المصروفات_المتكررة", "System_Metadata"}.issubset(wb.sheetnames)
        assert wb["البنود_المالية"].max_row == 2
        assert wb["System_Metadata"].sheet_state == "hidden"
        assert wb["System_Metadata"]["B2"].value.startswith("hesba-")

def test_excel_backup_incremental_merge_keeps_excel_only_rows(app):
    with app.app_context():
        user = make_user()
        account = FinancialAccount(name="Home", initial_amount=100)
        db.session.add(account)
        db.session.flush()
        db.session.add(Expense(account_id=account.id, amount=10, description="Food"))
        db.session.commit()
        first = merge_workbook(None, user.id)
        wb = load_workbook(BytesIO(first))
        wb["تفاصيل_المصروفات"].append([999, account.id, 25, "Excel-only historical row", "2026-01-01", "2026-01-01", "2026-01-01", None])
        buf = BytesIO()
        wb.save(buf)
        db.session.add(Expense(account_id=account.id, amount=15, description="Transport"))
        db.session.commit()
        second = merge_workbook(buf.getvalue(), user.id)
        result = load_workbook(BytesIO(second))
        rows = list(result["تفاصيل_المصروفات"].values)
        ids = [row[0] for row in rows[1:]]
        assert 999 in ids
        assert len(ids) == 3

def test_accounts_sheet_exports_total_funding_not_only_initial_amount(app):
    from app.models import FundingTransaction
    with app.app_context():
        user = make_user()
        account = FinancialAccount(name="Car profit", initial_amount=10000)
        db.session.add(account)
        db.session.flush()
        db.session.add_all([
            FundingTransaction(account_id=account.id, amount=10000, description="Second amount"),
        ])
        db.session.commit()
        content = merge_workbook(None, user.id)
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb["البنود_المالية"]
        headers = [cell.value for cell in ws[1]]
        total_col = headers.index("إجمالي_الأموال_المضافة") + 1
        assert ws.cell(row=2, column=total_col).value == 20000


def test_excel_backup_can_update_same_users_workbook(app):
    with app.app_context():
        user = make_user()
        account = FinancialAccount(name="Home", initial_amount=100)
        db.session.add(account)
        db.session.commit()

        first = merge_workbook(None, user.id)
        db.session.add(Expense(account_id=account.id, amount=25, description="Groceries"))
        db.session.commit()

        updated = merge_workbook(first, user.id)
        wb = load_workbook(BytesIO(updated))
        assert wb["تفاصيل_المصروفات"].max_row == 2
        assert wb["تفاصيل_المصروفات"]["D2"].value == "Groceries"


def test_excel_backup_rejects_another_users_workbook(app):
    import pytest
    with app.app_context():
        owner = make_user()
        other = User(username="other")
        other.set_password("another-password")
        db.session.add(other)
        db.session.commit()

        first = merge_workbook(None, owner.id)
        with pytest.raises(ValueError, match="حساب مختلف"):
            merge_workbook(first, other.id)
